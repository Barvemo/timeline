#!/usr/bin/env python3

import os
import sys
import glob
import argparse
from datetime import datetime
import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

'''
TODO: general cleanup, please...
TODO: Add logger, log file should be located in output folder with the resulting timeline.xlxs
TODO: split file into a main file, utils-file etc.
TODO: add readme
'''

def parse_args():
    parser = argparse.ArgumentParser(description="Merge CSV logs into an Excel timeline")
    parser.add_argument("-i", "--input", required=True, help="Input folder with CSV files")
    parser.add_argument("-o", "--output", required=False, help="Output folder for end product")
    return parser.parse_args()

def build_output_file(input_path, output_arg=None):
    if output_arg:
        output_folder = os.path.join(output_arg, f"timeline_{datetime.now().strftime('%Y-%m-%d_%H-%M-%S_%f')}")
    else:
        output_folder = os.path.join(input_path, f"timeline_{datetime.now().strftime('%Y-%m-%d_%H-%M-%S_%f')}")

    os.makedirs(output_folder, exist_ok=True)

    output_file = os.path.join(output_folder, "timeline.xlsx")
    return output_file

def get_csv_files(input_path):
    files = sorted(glob.glob(os.path.join(input_path, "*.csv")))
    if not files:
        sys.exit(1)
    return files

def process_files(files, output_file):
    dfs = []
    skipped_files = []

    for f in files:
        if os.path.getsize(f) == 0:
            skipped_files.append(f)
            continue

        try:
            df = pd.read_csv(f, sep=None, engine="python", encoding="utf-8")
            dfs.append(df)
        except Exception as e:
            skipped_files.append(f)

    if not dfs:
        sys.exit(1)

    # Merge all
    merged = pd.concat(dfs, ignore_index=True)

    # Parse Timestamp column if present
    if "Timestamp" in merged.columns:
        merged["Timestamp"] = pd.to_datetime(merged["Timestamp"], errors="coerce")
        merged = merged.sort_values(by="Timestamp")

    # Export to Excel
    merged.to_excel(output_file, index=False)

    # Enhance Excel file
    enhance_excel(output_file)

def enhance_excel(output_file):
    wb = load_workbook(output_file)
    ws = wb.active

    table = Table(displayName="LogTable", ref=ws.dimensions)
    style = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False
    )
    table.tableStyleInfo = style
    ws.add_table(table)

    ws.freeze_panes = "A2"

    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass
        ws.column_dimensions[col_letter].width = min(max_length + 2, 50)

    wb.save(output_file)

def main():
    args = parse_args()
    input_path = os.path.abspath(args.input)

    files = get_csv_files(input_path)
    output_file = build_output_file(input_path, args.output)
    process_files(files, output_file)

if __name__ == "__main__":
    main()